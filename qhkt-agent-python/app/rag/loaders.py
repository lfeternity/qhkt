from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.rag.transcript import encode_segments, parse_subtitle

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".html", ".htm", ".srt", ".vtt"}


def load_documents(
    filename: str,
    content: bytes,
    *,
    metadata: dict[str, Any] | None = None,
) -> list[Document]:
    """Load supported course sources into LangChain Documents with provenance."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"不支持的知识文件类型: {suffix or 'unknown'}")
    meta = {**(metadata or {}), "source_name": filename, "file_extension": suffix}
    if suffix == ".pdf":
        return _pdf_documents(content, meta)
    elif suffix == ".docx":
        text = _docx_text(content)
    elif suffix in {".html", ".htm"}:
        text = _html_text(content)
    elif suffix in {".srt", ".vtt"}:
        segments = parse_subtitle(filename, content.decode("utf-8-sig"))
        return [Document(page_content=encode_segments(segments), metadata={**meta, "timeline": True})]
    else:
        text = content.decode("utf-8-sig")
    text = re.sub(r"\x00", " ", text).strip()
    if not text:
        raise ValueError("知识文件内容为空")
    return [Document(page_content=text, metadata=meta)]


def split_documents(documents: list[Document], *, chunk_size: int = 1200, chunk_overlap: int = 160) -> list[Document]:
    if chunk_size <= chunk_overlap:
        raise ValueError("chunk_size 必须大于 chunk_overlap")
    splitter = RecursiveCharacterTextSplitter(chunk_size=chunk_size, chunk_overlap=chunk_overlap, separators=["\n\n", "\n", "。", "！", "？", ".", " ", ""])
    return splitter.split_documents(documents)


def _pdf_text(content: bytes) -> str:
    try:
        import fitz
    except ImportError as error:  # pragma: no cover - optional production dependency
        raise RuntimeError("PDF Loader 需要安装 pymupdf") from error
    pages: list[str] = []
    with fitz.open(stream=content, filetype="pdf") as document:
        for index, page in enumerate(document):
            pages.append(f"\n[page:{index + 1}]\n{page.get_text()}")
    return "\n".join(pages)


def _pdf_documents(content: bytes, metadata: dict[str, Any]) -> list[Document]:
    try:
        import fitz
    except ImportError as error:  # pragma: no cover - optional production dependency
        raise RuntimeError("PDF Loader 需要安装 pymupdf") from error
    documents: list[Document] = []
    with fitz.open(stream=content, filetype="pdf") as source:
        for index, page in enumerate(source, start=1):
            text = page.get_text().strip()
            if text:
                documents.append(Document(page_content=text, metadata={**metadata, "page": index}))
    if not documents:
        raise ValueError("知识文件内容为空")
    return documents


def _docx_text(content: bytes) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as error:  # pragma: no cover - optional production dependency
        raise RuntimeError("DOCX Loader 需要安装 python-docx") from error
    document = DocxDocument(io.BytesIO(content))
    return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())


def _html_text(content: bytes) -> str:
    try:
        from bs4 import BeautifulSoup
    except ImportError as error:  # pragma: no cover - optional production dependency
        raise RuntimeError("HTML Loader 需要安装 beautifulsoup4") from error
    soup = BeautifulSoup(content.decode("utf-8-sig"), "html.parser")
    for node in soup(["script", "style", "noscript"]):
        node.decompose()
    return soup.get_text("\n")
